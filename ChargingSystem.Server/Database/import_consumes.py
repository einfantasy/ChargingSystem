import openpyxl
import sqlite3
import sys
from uuid import uuid4

xlsx_path = r'C:\Users\einfa\source\repos\ChargingSystem\ChargingSystem.Server\Database\data_processed.xlsx'
db_path = r'C:\Users\einfa\source\repos\ChargingSystem\ChargingSystem.Server\Database\app.db'


def progress(current, total, prefix='', length=40):
    pct = current / total if total > 0 else 1
    filled = int(length * pct)
    bar = '\u2588' * filled + '\u2591' * (length - filled)
    sys.stdout.write(f'\r{prefix}|{bar}| {pct*100:5.1f}% ({current}/{total})')
    sys.stdout.flush()
    if current >= total:
        print()


wb = openpyxl.load_workbook(xlsx_path, data_only=True)
ws = wb.active

conn = sqlite3.connect(db_path)
cur = conn.cursor()

cur.execute('SELECT StudentId, Name FROM Students')
student_map = {row[1]: row[0] for row in cur.fetchall()}

cur.execute('SELECT TeacherId, Subject FROM Teachers')
teacher_map = {}
for row in cur.fetchall():
    if row[1] not in teacher_map:
        teacher_map[row[1]] = row[0]


def get_student_info(row_num):
    student_name = ws.cell(row=row_num, column=1).value
    if student_name:
        return student_name

    return None


def count_total():
    total = 0
    prev = None
    for row_num in range(4, 59):
        name = ws.cell(row=row_num, column=1).value
        if name:
            prev = name
        if not prev or prev not in student_map:
            continue
        rp = ws.cell(row=row_num, column=10).value
        rh = ws.cell(row=row_num, column=3).value
        if rp is None or rh is None:
            continue
        for col in range(12, ws.max_column + 1):
            if ws.cell(row=row_num, column=col).value is not None:
                subject = ws.cell(row=2, column=col).value
                date_val = ws.cell(row=3, column=col).value
                if date_val is not None and teacher_map.get(subject):
                    total += 1
    return total


total_records = count_total()
print(f'Total records to insert: {total_records}\n')

inserted = 0
skipped_no_student = 0
skipped_no_recharge = 0
prev_student_name = None

for row_num in range(4, 59):
    student_name = ws.cell(row=row_num, column=1).value
    if student_name:
        prev_student_name = student_name

    if not prev_student_name or prev_student_name not in student_map:
        skipped_no_student += 1
        continue

    student_id = student_map[prev_student_name]

    recharge_price = ws.cell(row=row_num, column=10).value
    if recharge_price is None:
        skipped_no_recharge += 1
        continue
    recharge_hour = ws.cell(row=row_num, column=3).value
    gift_hour = ws.cell(row=row_num, column=4).value
    if recharge_hour is None:
        skipped_no_recharge += 1
        continue
    from decimal import Decimal
    recharge_hour = str(Decimal(recharge_hour) + Decimal(gift_hour))
    
    cost = ws.cell(row=row_num, column=11).value
    if cost is None:
        cost = 0

    sql_query = f'SELECT Id FROM Recharges WHERE StudentId = "{student_id}" AND RechargePrice = {recharge_price} AND RechargeHour = {recharge_hour} AND Cost = {cost}'
    cur.execute(sql_query)
    recharge_row = cur.fetchone()
    if not recharge_row:
        #print(f'\nWARNING: Row {row_num} student "{prev_student_name}" - no matching Recharge (price={recharge_price}, hour={recharge_hour}, cost={cost})')
        #skipped_no_recharge += 1
        sql_query = f'SELECT Id FROM Recharges WHERE StudentId = "{student_id}" AND RechargeHour = {recharge_hour} AND Cost = {cost}'
        cur.execute(sql_query)
        recharge_row = cur.fetchone()
        if not recharge_row:
            print(f'\nWARNING: Row {row_num} student "{prev_student_name}" - no matching Recharge (price={recharge_price}, hour={recharge_hour}, cost={cost})')
            skipped_no_recharge += 1
            continue

    recharge_id = recharge_row[0]

    for col in range(12, ws.max_column + 1):
        hour_val = ws.cell(row=row_num, column=col).value
        if hour_val is None:
            continue

        subject = ws.cell(row=2, column=col).value
        date_val = ws.cell(row=3, column=col).value

        from datetime import datetime
        dt_obj = datetime.strptime(date_val, '%Y年%m月%d日')
        formatted_date = dt_obj.strftime('%Y-%m-%d 00:00:00')

        if formatted_date is None:
            continue

        date_str = formatted_date

        teacher_id = teacher_map.get(subject)
        if teacher_id is None:
            print(f'\nWARNING: Row {row_num}, Col {col} - subject "{subject}" not found in Teachers')
            continue

        if subject == '英语':
            cost_at_time = 0
        else:
            cost_at_time = cost

        consume_id = str(uuid4()).upper()

        sql_query = f"INSERT INTO Consumes (Id, RechargeId, TeacherId, Hour, Date, CostAtTime) VALUES ('{consume_id}', '{recharge_id}', '{teacher_id}', {hour_val}, '{date_str}', {cost_at_time})"
        cur.execute(sql_query)
        inserted += 1
        progress(inserted, total_records, prefix='  Insert: ')

conn.commit()
conn.close()
wb.close()

print(f'Done. Inserted: {inserted}, Skipped(no student): {skipped_no_student}, Skipped(no recharge): {skipped_no_recharge}')
